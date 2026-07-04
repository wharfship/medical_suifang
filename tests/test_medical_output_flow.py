import os
import tempfile
import unittest

from openpyxl import load_workbook

import medical_output_flow
from medical_output_flow import (
    build_followup_result_dirname,
    build_patient_storage_name,
    ensure_patient_output_dir,
    persist_followup_export,
    update_patient_summary_workbook,
)


PATIENT_NAME = "李同学"
STUDENT_ID = "30291834"
FOLLOWUP_DATE = "2025.06.12"
FOLLOWUP_LABEL = "第一次随访2026.02.07-2026.02.18"


class MedicalOutputPersistenceTests(unittest.TestCase):
    def test_build_followup_result_dirname_uses_readable_suffix(self):
        self.assertEqual(build_followup_result_dirname(FOLLOWUP_DATE), "2025.06.12随访结果")

    def test_build_followup_result_dirname_prefers_custom_followup_label(self):
        self.assertEqual(build_followup_result_dirname(FOLLOWUP_DATE, followup_label=FOLLOWUP_LABEL), FOLLOWUP_LABEL)

    def test_build_followup_result_dirname_sanitizes_time_style_label_for_windows(self):
        self.assertEqual(
            build_followup_result_dirname(followup_label="2026/06/29/18:24随访"),
            "2026／06／29／18：24随访",
        )

    def test_persist_followup_export_saves_excel_and_uploaded_report_in_patient_folder(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_excel_path = os.path.join(temp_dir, "medical_data.xlsx")
            uploaded_report_path = os.path.join(temp_dir, "report.pdf")
            output_dir = os.path.join(temp_dir, "outputs")

            with open(source_excel_path, "wb") as handle:
                handle.write(b"fake-excel")
            with open(uploaded_report_path, "wb") as handle:
                handle.write(b"fake-report")

            output_path = persist_followup_export(
                source_excel_path,
                output_dir=output_dir,
                uploaded_report_path=uploaded_report_path,
                patient_name=PATIENT_NAME,
                student_id=STUDENT_ID,
                followup_date=FOLLOWUP_DATE,
                followup_label=FOLLOWUP_LABEL,
            )

            submission_dir = os.path.dirname(output_path)
            copied_report_path = os.path.join(submission_dir, "report.pdf")
            expected_dir = os.path.join(
                output_dir,
                build_patient_storage_name(PATIENT_NAME, STUDENT_ID),
                FOLLOWUP_LABEL,
            )

            self.assertTrue(os.path.exists(output_path))
            self.assertTrue(os.path.isdir(submission_dir))
            self.assertEqual(submission_dir, expected_dir)
            self.assertEqual(os.path.basename(output_path), "medical_data.xlsx")
            self.assertTrue(os.path.exists(copied_report_path))
            with open(copied_report_path, "rb") as handle:
                self.assertEqual(handle.read(), b"fake-report")

    def test_persist_followup_export_ignores_missing_uploaded_report_path(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_excel_path = os.path.join(temp_dir, "medical_data.xlsx")
            output_dir = os.path.join(temp_dir, "outputs")

            with open(source_excel_path, "wb") as handle:
                handle.write(b"fake-excel")

            output_path = persist_followup_export(
                source_excel_path,
                output_dir=output_dir,
                uploaded_report_path=os.path.join(temp_dir, "missing.xlsx"),
                patient_name=PATIENT_NAME,
                student_id=STUDENT_ID,
                followup_date=FOLLOWUP_DATE,
                followup_label=FOLLOWUP_LABEL,
            )

            self.assertTrue(os.path.exists(output_path))
            self.assertEqual(len(os.listdir(os.path.dirname(output_path))), 1)

    def test_ensure_patient_output_dir_uses_patient_name(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_dir = os.path.join(temp_dir, "outputs")

            patient_output_dir = ensure_patient_output_dir(
                output_dir=output_dir,
                patient_name=PATIENT_NAME,
                student_id=STUDENT_ID,
                followup_date=FOLLOWUP_DATE,
                followup_label=FOLLOWUP_LABEL,
            )

            self.assertEqual(
                str(patient_output_dir),
                os.path.join(
                    output_dir,
                    build_patient_storage_name(PATIENT_NAME, STUDENT_ID),
                    FOLLOWUP_LABEL,
                ),
            )
            self.assertTrue(os.path.isdir(str(patient_output_dir)))

    def test_persist_followup_export_overwrites_existing_files_in_patient_folder(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_excel_path = os.path.join(temp_dir, "medical_data.xlsx")
            uploaded_report_path = os.path.join(temp_dir, "report.pdf")
            output_dir = os.path.join(temp_dir, "outputs")
            patient_output_dir = os.path.join(
                output_dir,
                build_patient_storage_name(PATIENT_NAME, STUDENT_ID),
                FOLLOWUP_LABEL,
            )

            with open(source_excel_path, "wb") as handle:
                handle.write(b"first-excel")
            with open(uploaded_report_path, "wb") as handle:
                handle.write(b"first-report")

            first_output_path = persist_followup_export(
                source_excel_path,
                output_dir=output_dir,
                uploaded_report_path=uploaded_report_path,
                patient_name=PATIENT_NAME,
                student_id=STUDENT_ID,
                followup_date=FOLLOWUP_DATE,
                followup_label=FOLLOWUP_LABEL,
            )

            with open(source_excel_path, "wb") as handle:
                handle.write(b"second-excel")
            with open(uploaded_report_path, "wb") as handle:
                handle.write(b"second-report")

            second_output_path = persist_followup_export(
                source_excel_path,
                output_dir=output_dir,
                uploaded_report_path=uploaded_report_path,
                patient_name=PATIENT_NAME,
                student_id=STUDENT_ID,
                followup_date=FOLLOWUP_DATE,
                followup_label=FOLLOWUP_LABEL,
            )

            self.assertEqual(os.path.dirname(first_output_path), patient_output_dir)
            self.assertEqual(os.path.dirname(second_output_path), patient_output_dir)
            with open(second_output_path, "rb") as handle:
                self.assertEqual(handle.read(), b"second-excel")
            with open(os.path.join(patient_output_dir, "report.pdf"), "rb") as handle:
                self.assertEqual(handle.read(), b"second-report")

    def test_update_patient_summary_workbook_builds_three_sheets(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            patient_dir = os.path.join(
                temp_dir,
                build_patient_storage_name(PATIENT_NAME, STUDENT_ID),
            )
            first_followup_dir = os.path.join(patient_dir, "第一次随访2026.02.07-2026.02.18")
            second_followup_dir = os.path.join(patient_dir, "第二次随访2026.02.19-2026.02.28")
            os.makedirs(first_followup_dir, exist_ok=True)
            os.makedirs(second_followup_dir, exist_ok=True)

            import pandas as pd

            pd.DataFrame(
                [
                    {"填写内容": "当前有无高血压", "填写数据": "无", "数据原始依据": "patient: 没有高血压", "解析状态": "done"},
                ]
            ).to_excel(os.path.join(first_followup_dir, "medical_data.xlsx"), index=False)
            pd.DataFrame(
                [
                    {"项目名称": "肌酐（酶法）", "结果": "117.8", "单位": "umol/L"},
                ]
            ).to_excel(os.path.join(first_followup_dir, "血生化_血清肌酐_提取结果.xlsx"), index=False)
            pd.DataFrame(
                [
                    {"项目名称": "尿蛋白", "结果": "-", "单位": ""},
                ]
            ).to_excel(os.path.join(second_followup_dir, "尿常规_尿蛋白_尿潜血_提取结果.xlsx"), index=False)

            summary_path = update_patient_summary_workbook(
                output_dir=temp_dir,
                patient_name=PATIENT_NAME,
                student_id=STUDENT_ID,
            )

            self.assertTrue(os.path.exists(summary_path))
            workbook = load_workbook(summary_path, data_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["medical_data", "血生化_血清肌酐_提取结果", "尿常规_尿蛋白_尿潜血_提取结果"],
            )
            medical_sheet = workbook["medical_data"]
            self.assertEqual(medical_sheet.cell(1, 1).value, "第一次随访2026.02.07-2026.02.18")
            self.assertEqual(medical_sheet.cell(2, 1).value, "填写内容")
            self.assertEqual(medical_sheet.cell(2, 2).value, "填写数据")
            self.assertEqual(medical_sheet.cell(2, 3).value, "数据原始依据")
            self.assertEqual(medical_sheet.cell(3, 1).value, "当前有无高血压")
            self.assertEqual(medical_sheet.cell(1, 5).value, "第二次随访2026.02.19-2026.02.28")

            blood_sheet = workbook["血生化_血清肌酐_提取结果"]
            self.assertEqual(blood_sheet.cell(1, 1).value, "第一次随访2026.02.07-2026.02.18")
            self.assertEqual(blood_sheet.cell(3, 1).value, "肌酐（酶法）")
            self.assertEqual(blood_sheet.cell(1, 5).value, "第二次随访2026.02.19-2026.02.28")

            urine_sheet = workbook["尿常规_尿蛋白_尿潜血_提取结果"]
            self.assertEqual(urine_sheet.cell(1, 1).value, "第一次随访2026.02.07-2026.02.18")
            self.assertEqual(urine_sheet.cell(1, 5).value, "第二次随访2026.02.19-2026.02.28")
            self.assertEqual(urine_sheet.cell(3, 5).value, "尿蛋白")
            self.assertFalse(os.path.exists(os.path.join(patient_dir, ".汇总结果.xlsx.tmp")))

    def test_update_patient_summary_workbook_writes_latest_copy_when_primary_is_locked(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            patient_dir = os.path.join(
                temp_dir,
                build_patient_storage_name(PATIENT_NAME, STUDENT_ID),
            )
            followup_dir = os.path.join(patient_dir, "第一次随访2026.02.07-2026.02.18")
            os.makedirs(followup_dir, exist_ok=True)

            import pandas as pd

            pd.DataFrame(
                [
                    {"填写内容": "当前有无高血压", "填写数据": "无", "数据原始依据": "patient: 没有高血压"},
                ]
            ).to_excel(os.path.join(followup_dir, "medical_data.xlsx"), index=False)

            real_replace = os.replace

            def fake_replace(src, dst):
                if str(dst).endswith("汇总结果.xlsx"):
                    raise PermissionError("locked")
                return real_replace(src, dst)

            with unittest.mock.patch.object(medical_output_flow.os, "replace", side_effect=fake_replace):
                summary_path = update_patient_summary_workbook(
                    output_dir=temp_dir,
                    patient_name=PATIENT_NAME,
                    student_id=STUDENT_ID,
                )

            self.assertTrue(summary_path.endswith("汇总结果_最新副本.xlsx"))
            self.assertTrue(os.path.exists(summary_path))


if __name__ == "__main__":
    unittest.main()
