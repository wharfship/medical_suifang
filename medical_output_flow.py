import os
import shutil
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "outputs"
DEFAULT_PATIENT_NAME = "\u674e\u540c\u5b66"
DEFAULT_STUDENT_ID = ""
SUMMARY_WORKBOOK_NAME = "汇总结果.xlsx"
SUMMARY_WORKBOOK_LATEST_COPY_NAME = "汇总结果_最新副本.xlsx"
SUMMARY_SHEET_SPECS = [
    {
        "sheet_name": "medical_data",
        "source_filename": "medical_data.xlsx",
        "columns": ["填写内容", "填写数据", "数据原始依据"],
    },
    {
        "sheet_name": "血生化_血清肌酐_提取结果",
        "source_filename": "血生化_血清肌酐_提取结果.xlsx",
        "columns": ["项目名称", "结果", "单位"],
    },
    {
        "sheet_name": "尿常规_尿蛋白_尿潜血_提取结果",
        "source_filename": "尿常规_尿蛋白_尿潜血_提取结果.xlsx",
        "columns": ["项目名称", "结果", "单位"],
    },
]


def sanitize_storage_fragment(value, fallback="patient"):
    text = str(value or "").strip()
    for char in '\\/:*?"<>|':
        text = text.replace(char, "_")
    text = "_".join(part for part in text.split() if part)
    return text or fallback


def build_patient_storage_name(patient_name=DEFAULT_PATIENT_NAME, student_id=DEFAULT_STUDENT_ID):
    normalized_patient_name = sanitize_storage_fragment(patient_name, fallback="patient")
    normalized_student_id = sanitize_storage_fragment(student_id, fallback="unknown")
    if normalized_student_id and normalized_student_id != "unknown":
        return f"{normalized_patient_name}_{normalized_student_id}"
    return normalized_patient_name


def build_followup_result_dirname(followup_date=None, followup_label=None):
    label_text = str(followup_label or "").strip()
    if label_text:
        return (
            label_text
            .replace("\\", "／")
            .replace("/", "／")
            .replace(":", "：")
        )

    target_date = followup_date or date.today()
    if isinstance(target_date, str):
        date_text = target_date.strip()
    else:
        date_text = target_date.strftime("%Y.%m.%d")
    return f"{date_text}随访结果"


def _copy_session_artifacts(source_artifact_dir, target_dir, primary_file=None):
    if not source_artifact_dir:
        return
    artifact_dir = Path(source_artifact_dir)
    if not artifact_dir.exists() or not artifact_dir.is_dir():
        return

    primary_resolved = Path(primary_file).resolve() if primary_file else None
    target_dir = Path(target_dir)
    for artifact_path in artifact_dir.iterdir():
        if not artifact_path.is_file():
            continue
        if primary_resolved and artifact_path.resolve() == primary_resolved:
            continue
        shutil.copy2(artifact_path, target_dir / artifact_path.name)


def _load_first_sheet_rows(file_path):
    workbook = load_workbook(file_path, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    if not rows:
        return []

    max_used_width = 0
    for row in rows:
        last_used_index = -1
        for index, value in enumerate(row):
            if value not in (None, ""):
                last_used_index = index
        if last_used_index >= 0:
            max_used_width = max(max_used_width, last_used_index + 1)

    if max_used_width == 0:
        return []

    trimmed_rows = [row[:max_used_width] for row in rows]
    while trimmed_rows and all(value in (None, "") for value in trimmed_rows[-1]):
        trimmed_rows.pop()
    return trimmed_rows


def _project_table_rows(rows, expected_columns):
    if not expected_columns:
        return rows if rows else []
    if not rows:
        return [list(expected_columns)]

    header = [str(value).strip() if value is not None else "" for value in rows[0]]
    column_indexes = {column_name: header.index(column_name) for column_name in expected_columns if column_name in header}

    projected_rows = [list(expected_columns)]
    for row in rows[1:]:
        projected_rows.append(
            [
                row[column_indexes[column_name]] if column_name in column_indexes and column_indexes[column_name] < len(row) else None
                for column_name in expected_columns
            ]
        )
    return projected_rows


def _write_summary_sheet(worksheet, followup_dirs, source_filename, expected_columns):
    start_col = 1
    block_headers = list(expected_columns or [])
    for followup_dir in followup_dirs:
        source_path = Path(followup_dir) / source_filename
        rows = _project_table_rows(_load_first_sheet_rows(source_path) if source_path.exists() else [], block_headers)
        if not rows:
            rows = [block_headers]

        block_width = max(1, len(rows[0]))
        worksheet.cell(row=1, column=start_col, value=Path(followup_dir).name)
        for row_index, row_values in enumerate(rows, start=2):
            for col_offset, value in enumerate(row_values):
                worksheet.cell(row=row_index, column=start_col + col_offset, value=value)
        start_col += block_width + 1


def update_patient_summary_workbook(output_dir=OUTPUT_DIR, patient_name=DEFAULT_PATIENT_NAME, student_id=DEFAULT_STUDENT_ID):
    patient_dir = Path(output_dir) / build_patient_storage_name(patient_name, student_id)
    if not patient_dir.exists():
        return None

    followup_dirs = sorted(
        [path for path in patient_dir.iterdir() if path.is_dir()],
        key=lambda path: path.stat().st_mtime,
    )
    if not followup_dirs:
        return None

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    for spec in SUMMARY_SHEET_SPECS:
        worksheet = workbook.create_sheet(title=spec["sheet_name"])
        _write_summary_sheet(
            worksheet,
            followup_dirs,
            spec["source_filename"],
            spec["columns"],
        )

    summary_path = patient_dir / SUMMARY_WORKBOOK_NAME
    temp_summary_path = patient_dir / f".{SUMMARY_WORKBOOK_NAME}.tmp"
    latest_copy_path = patient_dir / SUMMARY_WORKBOOK_LATEST_COPY_NAME
    workbook.save(temp_summary_path)
    try:
        os.replace(temp_summary_path, summary_path)
        return str(summary_path)
    except OSError:
        os.replace(temp_summary_path, latest_copy_path)
        return str(latest_copy_path)


def ensure_patient_output_dir(
    output_dir=OUTPUT_DIR,
    patient_name=DEFAULT_PATIENT_NAME,
    patient_output_dir=None,
    student_id=DEFAULT_STUDENT_ID,
    followup_date=None,
    followup_label=None,
):
    if patient_output_dir:
        target_dir = Path(patient_output_dir)
    else:
        patient_dir = Path(output_dir) / build_patient_storage_name(patient_name, student_id)
        target_dir = patient_dir / build_followup_result_dirname(followup_date, followup_label=followup_label)
    target_dir.mkdir(parents=True, exist_ok=True)
    return target_dir


def persist_followup_export(
    source_excel_path,
    output_dir=OUTPUT_DIR,
    uploaded_report_path=None,
    patient_name=DEFAULT_PATIENT_NAME,
    patient_output_dir=None,
    student_id=DEFAULT_STUDENT_ID,
    followup_date=None,
    followup_label=None,
    source_artifact_dir=None,
):
    source_excel = Path(source_excel_path)
    submission_dir = ensure_patient_output_dir(
        output_dir=output_dir,
        patient_name=patient_name,
        patient_output_dir=patient_output_dir,
        student_id=student_id,
        followup_date=followup_date,
        followup_label=followup_label,
    )

    target_excel_path = submission_dir / source_excel.name
    shutil.copy2(source_excel, target_excel_path)
    _copy_session_artifacts(source_artifact_dir, submission_dir, primary_file=source_excel)

    if uploaded_report_path:
        report_path = Path(uploaded_report_path)
        if report_path.exists():
            shutil.copy2(report_path, submission_dir / report_path.name)

    return str(target_excel_path)
